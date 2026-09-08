#!/usr/bin/env python3
"""Advisory presentation review of authored Project Mark files.

LLM assistance is permitted. This is not an AI detector, an authenticity
certificate or a platform gate. It never changes files. Keep metadata truthful
or blank; do not invent an author, approval, timestamp or business history.

    mark_realism_check.py PACKAGE [--authored memo.docx,decision.xlsx] [--strict]

Default exit: 0 after inspection, 2 for incomplete inspection. --strict also
returns 1 for review notes. Layout still needs rendering and human review.
"""
import argparse
from pathlib import Path
import re
import sys

SUPPORTED = {".docx", ".xlsx", ".txt", ".md"}


def prose_report(text):
    markers = re.findall(
        r"\[(?:insert|add|replace|your)\b[^\]\n]*\]|\blorem ipsum\b",
        text, re.I)
    return [f"Possible unfinished placeholder: {marker}" for marker in sorted(set(markers))]


def docx_report(path):
    import docx
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    parts.extend(cell.text for table in document.tables for row in table.rows
                 for cell in row.cells)
    return prose_report("\n".join(parts))


def xlsx_report(path):
    import openpyxl
    workbook = openpyxl.load_workbook(path)
    issues = []
    try:
        for sheet in workbook:
            populated_rows = 0
            clipped = []
            for row in sheet:
                if any(cell.value is not None for cell in row):
                    populated_rows += 1
                for cell in row:
                    value = cell.value
                    if cell.data_type == "e" or (
                            cell.data_type == "f" and "#REF!" in str(value)):
                        issues.append(f"{sheet.title}!{cell.coordinate}: stored error or broken reference")
                    if (isinstance(value, str) and cell.data_type != "f"
                            and len(value) > 60 and not cell.alignment.wrap_text
                            and not any(cell.coordinate in area for area in sheet.merged_cells.ranges)):
                        if len(clipped) < 5:
                            clipped.append(cell.coordinate)
            if populated_rows > 30 and not sheet.freeze_panes:
                issues.append(f"{sheet.title}: long table without frozen headings; review navigation")
            if clipped:
                issues.append(f"{sheet.title}: review potentially clipped text at {', '.join(clipped)}")
    finally:
        workbook.close()
    return issues


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("package_dir")
    parser.add_argument("--authored", help="comma-separated relative paths or basenames to inspect")
    parser.add_argument("--strict", action="store_true", help="return 1 when review notes exist")
    args = parser.parse_args(argv)
    root = Path(args.package_dir)
    if not root.is_dir():
        print(f"Inspection incomplete: not a directory: {root}")
        return 2
    requested = {part.strip() for part in args.authored.split(",") if part.strip()} if args.authored else None
    matched = set()
    findings, errors = {}, []
    inspected = 0
    for path in sorted(root.rglob("*")):
        relative = path.relative_to(root)
        if not path.is_file() or any(part.startswith(".") for part in relative.parts):
            continue
        names = {relative.as_posix(), path.name}
        if requested is not None:
            hits = requested & names
            if not hits:
                continue
            matched.update(hits)
        if path.suffix.lower() not in SUPPORTED:
            if requested is not None:
                errors.append(f"{relative}: unsupported file type")
            continue
        try:
            if path.suffix.lower() == ".xlsx":
                issues = xlsx_report(path)
            elif path.suffix.lower() == ".docx":
                issues = docx_report(path)
            else:
                issues = prose_report(path.read_text(encoding="utf-8"))
            inspected += 1
            if issues:
                findings[relative.as_posix()] = issues
        except Exception as exc:
            errors.append(f"{relative}: inspection failed: {exc}")
    if requested:
        errors.extend(f"{name}: requested file missing" for name in sorted(requested - matched))
    if inspected == 0 and not errors:
        errors.append("No supported files inspected")
    for name, notes in findings.items():
        print(name)
        for note in notes:
            print(f"  Review: {note}")
    for error in errors:
        print(f"Incomplete: {error}")
    print(f"Inspected {inspected} files; {sum(map(len, findings.values()))} review notes; {len(errors)} inspection errors.")
    print("LLM assistance is allowed. Keep metadata truthful or blank. Regular layouts and single-sheet workbooks can be appropriate.")
    print("Render and review the files for usability, accurate figures, legible visuals and stakeholder relevance.")
    return 2 if errors else (1 if args.strict and findings else 0)


if __name__ == "__main__":
    sys.exit(main())
