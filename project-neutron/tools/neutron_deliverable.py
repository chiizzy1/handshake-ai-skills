#!/usr/bin/env python3
"""Golden workbook writing and rubric construction.

Both halves take their numbers from the same computed truth, so the deliverable
and the rubric cannot disagree. Hand-typing a figure into either one is the
defect this module exists to prevent.
"""
import decimal
import json

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

# Current rubric rules (the module, not the delivered examples, which still
# show +10/+8/+4/+2 and rubrics of 120 items).
ALLOWED_WEIGHTS = {1, 3, 5, 7, 9}
MIN_ITEMS, MAX_ITEMS = 20, 100
MAX_NEGATIVE_SHARE = 0.35
GOLDEN_PASS = 95.0

# Share of items each magnitude should occupy. All five bands are stated in
# the weighting module, not just the 9s — checking only the 9s let a rubric
# through with its 7s at 24 % against a 10-20 % target.
WEIGHT_SHARE = {
    9: (0.10, 0.15),
    7: (0.10, 0.20),
    5: (0.15, 0.25),
    3: (0.15, 0.25),
    1: (0.15, 0.20),
}


def r2(x):
    """Half-up to two decimals, halves away from zero."""
    return decimal.Decimal(str(x)).quantize(decimal.Decimal("0.01"),
                                            rounding=decimal.ROUND_HALF_UP)


def band(value, pct):
    """Tolerance band around a computed value, as (lo, hi) at 2 dp."""
    v = decimal.Decimal(str(value))
    p = decimal.Decimal(str(pct))
    return r2(v * (1 - p)), r2(v * (1 + p))


THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


class Workbook:
    """Writes the stacked, titled tables these take-offs are graded on."""

    def __init__(self, sheet_title, heading=None):
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = sheet_title
        self.row = 1
        if heading:
            self.ws.cell(row=1, column=1, value=heading).font = Font(bold=True,
                                                                     size=12)
            self.row = 3

    def table(self, title, columns, rows, total=None, widths=None):
        ws = self.ws
        ws.cell(row=self.row, column=1, value=title).font = Font(bold=True,
                                                                 size=11)
        self.row += 1
        for i, c in enumerate(columns, 1):
            cell = ws.cell(row=self.row, column=i, value=c)
            cell.font = Font(bold=True, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")
        self.row += 1
        for r in rows:
            self._write(r)
        if total is not None:
            self._write(total, bold=True)
        self.row += 1
        if widths:
            for col, w in zip("ABCDEFGHIJ", widths):
                ws.column_dimensions[col].width = w
        return self

    def _write(self, cells, bold=False):
        for i, v in enumerate(cells, 1):
            c = self.ws.cell(row=self.row, column=i)
            c.border = BORDER
            c.font = Font(bold=bold, size=10)
            if isinstance(v, decimal.Decimal):
                c.value = float(v)
                c.number_format = "0.00"
                c.alignment = Alignment(horizontal="right")
            elif isinstance(v, int):
                c.value = v
                c.number_format = "0"
                c.alignment = Alignment(horizontal="right")
            else:
                c.value = v
        self.row += 1

    def save(self, path, creator="Take-off", title=None):
        self.ws.freeze_panes = "A2"
        p = self.wb.properties
        p.creator = creator
        p.title = title or creator
        self.wb.save(path)
        return path


class Rubric:
    """Builds and validates a rubric against the current module rules."""

    def __init__(self):
        self.items = []

    def add(self, criterion, weight):
        if abs(weight) not in ALLOWED_WEIGHTS:
            raise ValueError(f"weight {weight} not in "
                             f"{sorted(ALLOWED_WEIGHTS)}: {criterion[:60]}")
        self.items.append({"n": len(self.items) + 1, "criterion": criterion,
                           "weight": weight})
        return self

    def positives(self):
        return sum(i["weight"] for i in self.items if i["weight"] > 0)

    def report(self, golden_misses=()):
        """`golden_misses` are criterion numbers the golden does not satisfy;
        normally empty, because the golden is built from the same truth."""
        n = len(self.items)
        neg = [i for i in self.items if i["weight"] < 0]
        pos = self.positives()
        lost = sum(i["weight"] for i in self.items
                   if i["n"] in golden_misses and i["weight"] > 0)
        pct = 100.0 * (pos - lost) / pos if pos else 0.0

        checks = {
            f"item count {n} in {MIN_ITEMS}-{MAX_ITEMS}":
                MIN_ITEMS <= n <= MAX_ITEMS,
            f"weights {sorted({abs(i['weight']) for i in self.items})} allowed":
                {abs(i["weight"]) for i in self.items} <= ALLOWED_WEIGHTS,
            f"negatives {len(neg)}/{n} = {100*len(neg)/n:.0f}% <= 35%":
                len(neg) / n <= MAX_NEGATIVE_SHARE,
        }
        for w in (9, 7, 5, 3, 1):
            k = sum(1 for i in self.items if abs(i["weight"]) == w)
            lo, hi = WEIGHT_SHARE[w]
            checks[f"{w}-weight {k}/{n} = {100*k/n:.0f}% in "
                   f"{lo*100:.0f}-{hi*100:.0f}%"] = lo <= k / n <= hi
        checks[f"golden scores {pct:.1f}% >= {GOLDEN_PASS}%"] = pct >= GOLDEN_PASS
        width = max(len(k) for k in checks)
        for k, ok in checks.items():
            print(f"  {k:<{width}}  {'PASS' if ok else 'FAIL'}")
        allok = all(checks.values())
        print(f"  {'RUBRIC VALID':<{width}}  {allok}")
        return allok

    def save(self, path):
        json.dump(self.items, open(path, "w"), indent=2)
        return path
